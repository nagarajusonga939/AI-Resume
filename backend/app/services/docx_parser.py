import os
import re
from docx import Document


def clean_docx_text(text: str) -> str:
    """Normalize extracted text whitespace and linebreaks."""
    if not text:
        return ""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'\n{3,}', '\n\n', text)
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n')]
    return '\n'.join(lines).strip()


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract plain text from DOCX files, including paragraphs and tables.
    
    Args:
        file_path (str): Path to the DOCX file.
        
    Returns:
        str: Clean combined text.
        
    Raises:
        FileNotFoundError: If the file does not exist.
        ValueError: If file is invalid or corrupted.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found at: {file_path}")

    try:
        doc = Document(file_path)
    except Exception as e:
        raise ValueError(f"Failed to open or decode DOCX file: {str(e)}")

    extracted_content = []

    # Extract all paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            extracted_content.append(p_text)

    # Extract all tables (often used in resumes for skills/contact layout)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                # Remove duplicate cell text (from merged cells)
                unique_cells = []
                for cell in row_cells:
                    if cell not in unique_cells:
                        unique_cells.append(cell)
                extracted_content.append(" | ".join(unique_cells))

    full_text = "\n".join(extracted_content)
    clean_text = clean_docx_text(full_text)

    if not clean_text:
        raise ValueError("The DOCX document is empty or contains no readable text.")

    return clean_text
